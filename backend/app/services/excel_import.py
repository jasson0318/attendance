"""
Excel 月曆式排班匯入。

盡量相容常見格式：
- 橫列：排班代碼／姓名
- 橫欄：日期 1～31 或完整日期
- 儲存格：08:00-17:00、0800-1700、休、假、OFF、空白

匯入流程：解析 → 預覽 → 異常檢查 → 確認匯入
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from io import BytesIO
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..timeutil import today_local

from ..models import Employee, Schedule

TIME_RANGE = re.compile(
    r"(?P<sh>\d{1,2}):?(?P<sm>\d{2})\s*[-~～到至]\s*(?P<eh>\d{1,2}):?(?P<em>\d{2})"
)
DAY_OFF_TOKENS = {"休", "假", "休假", "例休", "特休", "OFF", "off", "Off", "x", "X", "－", "-", "—"}


@dataclass
class ParsedCell:
    row: int
    col: int
    schedule_code: str
    work_date: date
    start: Optional[time] = None
    end: Optional[time] = None
    break_start: Optional[time] = None
    break_end: Optional[time] = None
    is_day_off: bool = False
    raw: str = ""


@dataclass
class ImportIssue:
    level: str  # error | warning
    code: str
    message: str
    row: Optional[int] = None
    col: Optional[int] = None
    schedule_code: Optional[str] = None
    work_date: Optional[str] = None


@dataclass
class ImportPreview:
    store_id: int
    year: int
    month: int
    rows: list[dict[str, Any]] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)
    ok_count: int = 0
    error_count: int = 0


def _parse_time_pair(text: str) -> Optional[tuple[time, time]]:
    m = TIME_RANGE.search(text.replace("：", ":"))
    if not m:
        return None
    sh, sm = int(m.group("sh")), int(m.group("sm"))
    eh, em = int(m.group("eh")), int(m.group("em"))
    if not (0 <= sh <= 23 and 0 <= eh <= 23 and 0 <= sm <= 59 and 0 <= em <= 59):
        return None
    return time(sh, sm), time(eh, em)


def _default_break(emp: Optional[Employee]) -> tuple[Optional[time], Optional[time]]:
    if emp and emp.break_start and emp.break_end:
        return emp.break_start, emp.break_end
    return time(12, 0), time(13, 0)


def _detect_date_headers(ws) -> dict[int, date]:
    """掃描前 5 列尋找日期欄。"""
    mapping: dict[int, date] = {}
    year_hint = None
    month_hint = None

    # 標題列可能有「2026年9月」
    for r in range(1, 6):
        for c in range(1, min(ws.max_column or 1, 40) + 1):
            val = ws.cell(r, c).value
            if val is None:
                continue
            s = str(val)
            ym = re.search(r"(20\d{2})\s*年\s*(\d{1,2})\s*月", s)
            if ym:
                year_hint = int(ym.group(1))
                month_hint = int(ym.group(2))
            if isinstance(val, datetime):
                mapping[c] = val.date()
            elif isinstance(val, date):
                mapping[c] = val

    if mapping:
        return mapping

    # 純數字日 1~31
    for r in range(1, 6):
        nums = []
        for c in range(1, min(ws.max_column or 1, 40) + 1):
            val = ws.cell(r, c).value
            if isinstance(val, (int, float)) and 1 <= int(val) <= 31:
                nums.append((c, int(val)))
            elif isinstance(val, str) and val.strip().isdigit() and 1 <= int(val.strip()) <= 31:
                nums.append((c, int(val.strip())))
        if len(nums) >= 20 or (len(nums) >= 7 and year_hint and month_hint):
            y = year_hint or today_local().year
            m = month_hint or today_local().month
            for c, day in nums:
                try:
                    mapping[c] = date(y, m, day)
                except ValueError:
                    pass
            if mapping:
                return mapping
    return mapping


def _detect_code_column(ws, date_cols: set[int]) -> int:
    """找排班代碼欄（通常在日期欄左側）。"""
    for c in range(1, 10):
        if c in date_cols:
            continue
        texts = []
        for r in range(2, min(ws.max_row or 2, 40) + 1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                texts.append(str(v).strip())
        if texts and all(len(t) <= 10 for t in texts[:15]):
            return c
    return 1


def parse_excel(
    content: bytes,
    store_id: int,
    db: Session,
    year: Optional[int] = None,
    month: Optional[int] = None,
) -> ImportPreview:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    date_headers = _detect_date_headers(ws)
    if not date_headers:
        preview = ImportPreview(store_id=store_id, year=year or 0, month=month or 0)
        preview.issues.append(
            ImportIssue(level="error", code="unrecognizable", message="無法辨識日期欄，請確認為月曆式排班表")
        )
        preview.error_count = 1
        return preview

    dates = list(date_headers.values())
    y = year or dates[0].year
    m = month or dates[0].month
    preview = ImportPreview(store_id=store_id, year=y, month=m)

    # 月份檢查
    for d in dates:
        if d.year != y or d.month != m:
            preview.issues.append(
                ImportIssue(
                    level="error",
                    code="month_mismatch",
                    message=f"日期 {d} 與指定月份 {y}-{m:02d} 不符",
                    work_date=str(d),
                )
            )

    code_col = _detect_code_column(ws, set(date_headers.keys()))
    employees = (
        db.query(Employee)
        .filter(Employee.store_id == store_id)
        .all()
    )
    by_code = {e.schedule_code: e for e in employees}
    seen: set[tuple[str, date]] = set()

    max_row = ws.max_row or 1
    for r in range(1, max_row + 1):
        code_val = ws.cell(r, code_col).value
        if code_val is None:
            continue
        code = str(code_val).strip()
        if not code or code in ("排班代碼", "代碼", "姓名", "員工", "名稱"):
            continue
        # 略過純日期標題列
        if code.isdigit() and 1 <= int(code) <= 31:
            continue

        for c, work_date in date_headers.items():
            raw_val = ws.cell(r, c).value
            if raw_val is None or str(raw_val).strip() == "":
                continue
            raw = str(raw_val).strip()

            cell = ParsedCell(row=r, col=c, schedule_code=code, work_date=work_date, raw=raw)
            issues_before = len(preview.issues)

            emp = by_code.get(code)
            if not emp:
                preview.issues.append(
                    ImportIssue(
                        level="error",
                        code="unknown_code",
                        message=f"不存在的排班代碼：{code}",
                        row=r,
                        col=c,
                        schedule_code=code,
                        work_date=str(work_date),
                    )
                )
            else:
                if emp.store_id != store_id:
                    preview.issues.append(
                        ImportIssue(
                            level="error",
                            code="wrong_store",
                            message=f"員工 {code} 不屬於該店",
                            row=r,
                            schedule_code=code,
                            work_date=str(work_date),
                        )
                    )
                if not emp.is_active or (emp.leave_date and emp.leave_date <= work_date):
                    preview.issues.append(
                        ImportIssue(
                            level="error",
                            code="inactive",
                            message=f"已離職員工：{code}",
                            row=r,
                            schedule_code=code,
                            work_date=str(work_date),
                        )
                    )

            key = (code, work_date)
            if key in seen:
                preview.issues.append(
                    ImportIssue(
                        level="error",
                        code="duplicate",
                        message=f"重複排班：{code} {work_date}",
                        row=r,
                        schedule_code=code,
                        work_date=str(work_date),
                    )
                )
            seen.add(key)

            if work_date.year != y or work_date.month != m:
                preview.issues.append(
                    ImportIssue(
                        level="error",
                        code="date_error",
                        message=f"日期錯誤：{work_date}",
                        work_date=str(work_date),
                    )
                )

            if raw in DAY_OFF_TOKENS:
                cell.is_day_off = True
            else:
                pair = _parse_time_pair(raw)
                if not pair:
                    preview.issues.append(
                        ImportIssue(
                            level="error",
                            code="time_format",
                            message=f"時間格式錯誤：{raw}",
                            row=r,
                            col=c,
                            schedule_code=code,
                            work_date=str(work_date),
                        )
                    )
                else:
                    cell.start, cell.end = pair
                    bs, be = _default_break(emp)
                    cell.break_start, cell.break_end = bs, be

            row_ok = len(preview.issues) == issues_before
            preview.rows.append(
                {
                    "row": r,
                    "col": c,
                    "schedule_code": code,
                    "employee_id": emp.id if emp else None,
                    "employee_name": emp.name if emp else None,
                    "work_date": str(work_date),
                    "start": cell.start.strftime("%H:%M") if cell.start else None,
                    "break_start": cell.break_start.strftime("%H:%M") if cell.break_start else None,
                    "break_end": cell.break_end.strftime("%H:%M") if cell.break_end else None,
                    "end": cell.end.strftime("%H:%M") if cell.end else None,
                    "is_day_off": cell.is_day_off,
                    "raw": raw,
                    "ok": row_ok and emp is not None,
                }
            )
            if row_ok and emp is not None:
                preview.ok_count += 1

    preview.error_count = sum(1 for i in preview.issues if i.level == "error")
    return preview


def commit_import(db: Session, preview: ImportPreview, replace_month: bool = True) -> int:
    """確認匯入。有 error 的列略過；ok 列寫入。"""
    if replace_month and preview.year and preview.month:
        # 刪除該店該月既有排班（僅限匯入成功的員工可選，這裡清該店整月再寫入）
        from calendar import monthrange

        last = monthrange(preview.year, preview.month)[1]
        start = date(preview.year, preview.month, 1)
        end = date(preview.year, preview.month, last)
        existing = (
            db.query(Schedule)
            .filter(
                Schedule.store_id == preview.store_id,
                Schedule.work_date >= start,
                Schedule.work_date <= end,
            )
            .all()
        )
        for e in existing:
            db.delete(e)
        db.flush()

    count = 0
    for row in preview.rows:
        if not row.get("ok"):
            continue
        emp_id = row["employee_id"]
        wd = date.fromisoformat(row["work_date"])
        sched = (
            db.query(Schedule)
            .filter(Schedule.employee_id == emp_id, Schedule.work_date == wd)
            .first()
        )
        if not sched:
            sched = Schedule(
                employee_id=emp_id,
                store_id=preview.store_id,
                work_date=wd,
            )
            db.add(sched)
        sched.store_id = preview.store_id
        sched.is_day_off = bool(row["is_day_off"])
        sched.shift_type = "day_off" if row["is_day_off"] else "full_day"
        sched.start_time = (
            datetime.strptime(row["start"], "%H:%M").time() if row.get("start") else None
        )
        sched.break_start = (
            datetime.strptime(row["break_start"], "%H:%M").time()
            if row.get("break_start")
            else None
        )
        sched.break_end = (
            datetime.strptime(row["break_end"], "%H:%M").time()
            if row.get("break_end")
            else None
        )
        sched.end_time = (
            datetime.strptime(row["end"], "%H:%M").time() if row.get("end") else None
        )
        sched.source = "excel"
        count += 1
    db.commit()
    return count


def preview_to_dict(preview: ImportPreview) -> dict:
    return {
        "store_id": preview.store_id,
        "year": preview.year,
        "month": preview.month,
        "ok_count": preview.ok_count,
        "error_count": preview.error_count,
        "rows": preview.rows,
        "issues": [
            {
                "level": i.level,
                "code": i.code,
                "message": i.message,
                "row": i.row,
                "col": i.col,
                "schedule_code": i.schedule_code,
                "work_date": i.work_date,
            }
            for i in preview.issues
        ],
    }
