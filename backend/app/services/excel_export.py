"""Excel 匯出：每日／月報／年度。"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from sqlalchemy.orm import Session

from ..models import AnnualStatistic, AttendanceDailySummary, Employee, Store


def _mins_to_hm(m: int) -> str:
    h, mm = divmod(int(m or 0), 60)
    return f"{h}:{mm:02d}"


def export_daily(db: Session, work_date: date, store_id: int | None = None) -> bytes:
    q = db.query(AttendanceDailySummary).filter(AttendanceDailySummary.work_date == work_date)
    if store_id:
        q = q.filter(AttendanceDailySummary.store_id == store_id)
    rows = q.all()
    wb = Workbook()
    ws = wb.active
    ws.title = "每日出勤"
    headers = [
        "日期",
        "員工",
        "店舖",
        "排班上班",
        "實際上班",
        "排班休息開始",
        "實際休息開始",
        "排班休息結束",
        "實際休息結束",
        "排班下班",
        "實際下班",
        "休息(分)",
        "實際工作(分)",
        "打卡跨度(分)",
        "遲到(分)",
        "早退(分)",
        "補打卡",
        "曠職",
        "出勤狀態",
        "無排班出勤原因",
        "異常",
    ]
    ws.append(headers)
    for s in rows:
        emp = db.query(Employee).filter(Employee.id == s.employee_id).first()
        store = db.query(Store).filter(Store.id == s.store_id).first()
        ws.append(
            [
                str(s.work_date),
                emp.name if emp else "",
                store.name if store else "",
                s.scheduled_start.strftime("%H:%M") if s.scheduled_start else "",
                s.actual_start.strftime("%H:%M") if s.actual_start else "",
                s.scheduled_break_start.strftime("%H:%M") if s.scheduled_break_start else "",
                s.actual_break_start.strftime("%H:%M") if s.actual_break_start else "",
                s.scheduled_break_end.strftime("%H:%M") if s.scheduled_break_end else "",
                s.actual_break_end.strftime("%H:%M") if s.actual_break_end else "",
                s.scheduled_end.strftime("%H:%M") if s.scheduled_end else "",
                s.actual_end.strftime("%H:%M") if s.actual_end else "",
                s.break_minutes,
                s.work_minutes,
                s.span_minutes,
                s.late_minutes,
                s.early_leave_minutes,
                s.makeup_count,
                "是" if s.is_absent else "否",
                s.attendance_status or "",
                s.unscheduled_reason or "",
                s.anomaly_notes or "",
            ]
        )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_monthly(db: Session, year: int, month: int, store_id: int | None = None) -> bytes:
    from calendar import monthrange

    last = monthrange(year, month)[1]
    start = date(year, month, 1)
    end = date(year, month, last)
    q = db.query(AttendanceDailySummary).filter(
        AttendanceDailySummary.work_date >= start,
        AttendanceDailySummary.work_date <= end,
    )
    if store_id:
        q = q.filter(AttendanceDailySummary.store_id == store_id)
    rows = q.order_by(AttendanceDailySummary.work_date, AttendanceDailySummary.employee_id).all()
    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}-{month:02d}月報"
    ws.append(
        [
            "日期",
            "員工",
            "店舖",
            "工作時數",
            "打卡跨度",
            "遲到",
            "早退",
            "補打卡",
            "曠職",
            "異常",
        ]
    )
    for s in rows:
        emp = db.query(Employee).filter(Employee.id == s.employee_id).first()
        store = db.query(Store).filter(Store.id == s.store_id).first()
        ws.append(
            [
                str(s.work_date),
                emp.name if emp else "",
                store.name if store else "",
                _mins_to_hm(s.work_minutes),
                _mins_to_hm(s.span_minutes),
                s.late_minutes,
                s.early_leave_minutes,
                s.makeup_count,
                "是" if s.is_absent else "否",
                s.anomaly_notes or "",
            ]
        )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_annual(db: Session, year: int, store_id: int | None = None) -> bytes:
    q = db.query(AnnualStatistic).filter(AnnualStatistic.year == year)
    if store_id:
        q = q.filter(AnnualStatistic.store_id == store_id)
    rows = q.all()
    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}年度報表"
    ws.append(
        [
            "員工",
            "店舖",
            "年度",
            "約定出勤總時數",
            "實際出勤總時數(不含加班)",
            "實際打卡總時數(不含加班)",
            "曠職時數",
            "遲到次數",
            "遲到總時數",
            "早退次數",
            "早退總時數",
            "補打卡總次數",
            "超過額度補打卡次數",
            "無排班出勤次數",
            "無排班出勤時數",
        ]
    )
    for a in rows:
        emp = db.query(Employee).filter(Employee.id == a.employee_id).first()
        store = db.query(Store).filter(Store.id == a.store_id).first()
        ws.append(
            [
                emp.name if emp else "",
                store.name if store else "",
                a.year,
                _mins_to_hm(a.agreed_work_minutes),
                _mins_to_hm(a.actual_work_minutes),
                _mins_to_hm(a.actual_span_minutes),
                _mins_to_hm(a.absent_minutes),
                a.late_count,
                _mins_to_hm(a.late_minutes),
                a.early_leave_count,
                _mins_to_hm(a.early_leave_minutes),
                a.makeup_total,
                a.makeup_over_quota,
                a.unscheduled_count,
                _mins_to_hm(a.unscheduled_work_minutes),
            ]
        )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
